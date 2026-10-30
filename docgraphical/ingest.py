# -*- coding: utf-8 -*-
"""Document ingest: pptx/docx/pdf -> text -> LLM summary -> md saved next to sources.

Flow (driven by the Explorer dir-node "New document" dialog):
    upload (copied into target dir) -> summarize (LLM preview) ->
    confirm-save (md lands in the SAME dir, header records each source's
    repo-relative path, owning repo incrementally indexed).

Rule (user-mandated): the generated md MUST live next to the uploaded files
AND MUST contain the files' relative paths (repo-relative, so docgraph links
resolve). The header is prepended server-side — never left to the LLM.
"""
import datetime
import os
import re
from typing import Any, Dict, List, Optional

MAX_TEXT_CHARS = 12000  # per-file truncation budget for the LLM prompt
ALLOWED_EXTS = {".pptx", ".ppt", ".docx", ".doc", ".pdf",
                ".txt", ".md", ".markdown", ".csv", ".html", ".htm",
                ".eml", ".jpg", ".jpeg", ".png", ".webp", ".bmp",
                ".tiff", ".tif"}
PARSEABLE_EXTS = {".pptx", ".docx", ".pdf"}
# Direct-to-vision: no text layer to extract, skip straight to eyes.
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff", ".tif"}
# Plain-text readable without special libs.
TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".html", ".htm", ".eml"}
CSV_MAX_ROWS = 200  # spreadsheet rows kept before truncation note
# Vision fallback ("open the LLM's eyes"): page/slide images for image-only docs.
VISION_MAX_PAGES = 12   # per file
VISION_MAX_TOTAL = 16   # per summarize call
VISION_MAX_DIM = 1280   # downscale longest edge
VISION_JPEG_Q = 70
VISION_MIN_EDGE = 250   # skip icons/logos smaller than this


def FnSanitizeFilename(name: str) -> str:
    """Unicode-safe basename (unlike werkzeug's ASCII-only secure_filename)."""
    base = os.path.basename((name or "").strip().replace("\\", "/"))
    base = re.sub(r"[\x00-\x1f\x7f]", "", base).strip().strip(".")
    base = re.sub(r'[<>:"|?*\u200b-\u200f]', "_", base)
    return base or "untitled"


def FnIngestDirOk(target_dir: str, roots: List[str]) -> bool:
    """Target must resolve inside one of the search roots (no arbitrary writes)."""
    try:
        t = os.path.abspath(target_dir)
    except Exception:
        return False
    for r in roots or []:
        try:
            a = os.path.abspath(r)
        except Exception:
            continue
        if t == a or t.startswith(a + os.sep):
            return True
    return False


def FnFindOwningRepo(target_dir: str, repos: Dict[str, str]) -> Optional[str]:
    """Longest-prefix match of target_dir against known repo paths."""
    try:
        t = os.path.abspath(target_dir)
    except Exception:
        return None
    best = None
    for _name, p in (repos or {}).items():
        try:
            a = os.path.abspath(p)
        except Exception:
            continue
        if t == a or t.startswith(a + os.sep):
            if best is None or len(a) > len(best):
                best = a
    return best


def FnExtractText(path: str) -> Dict[str, Any]:
    """Extract plain text: office/pdf + txt/md/csv/html/eml. Old .ppt/.doc
    binaries rejected; image formats skip text entirely (vision path)."""
    ext = os.path.splitext(path or "")[1].lower()
    if ext not in ALLOWED_EXTS:
        return {"ok": False, "error": "不支援的格式: %s" % ext}
    if ext in IMAGE_EXTS:
        return {"ok": False, "error": "__IMAGE__"}
    if ext not in PARSEABLE_EXTS and ext not in TEXT_EXTS:
        return {"ok": False,
                "error": "舊版二進位格式 (%s) 無法解析，請用 Office 另存為 .pptx/.docx 後再上傳" % ext}
    try:
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            chunks = []
            for i, slide in enumerate(prs.slides, 1):
                parts = ["[Slide %d]" % i]
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame:
                        t = shape.text.strip()
                        if t:
                            parts.append(t)
                    try:
                        if shape.has_table:
                            for row in shape.table.rows:
                                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                                if cells:
                                    parts.append(" | ".join(cells))
                    except Exception:
                        pass
                if len(parts) > 1:
                    chunks.append("\n".join(parts))
            text = "\n\n".join(chunks)
        elif ext == ".docx":
            import docx
            doc = docx.Document(path)
            chunks = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            for tb in doc.tables:
                for row in tb.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        chunks.append(" | ".join(cells))
            text = "\n".join(chunks)
        elif ext in (".txt", ".md", ".markdown"):
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        elif ext == ".csv":
            import csv as _csv
            _rows = []
            with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
                for i, row in enumerate(_csv.reader(f)):
                    if i >= CSV_MAX_ROWS:
                        _rows.append("…（僅取前 %d 列）" % CSV_MAX_ROWS)
                        break
                    _rows.append("| " + " | ".join((c or "").strip() for c in row) + " |")
            text = "\n".join(_rows)
        elif ext in (".html", ".htm"):
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                raw = f.read()
            try:
                from bs4 import BeautifulSoup as _Soup
                _soup = _Soup(raw, "html.parser")
                for _tag in _soup(["script", "style", "noscript"]):
                    _tag.decompose()
                _t = "\n".join(ln.strip() for ln in _soup.get_text(separator="\n").splitlines() if ln.strip())
                _title = (_soup.title.string.strip() if _soup.title and _soup.title.string else "")
                text = (("# " + _title + "\n\n") if _title else "") + _t
            except Exception:
                _t = re.sub(r"(?s)<script.*?</script>|<style.*?</style>|<[^>]+>", " ", raw)
                text = "\n".join(ln.strip() for ln in re.sub(r"[ \t]+", " ", _t).splitlines() if ln.strip())
        elif ext == ".eml":
            import email as _email
            import email.policy as _epolicy
            with open(path, "rb") as f:
                _msg = _email.message_from_binary_file(f, policy=_epolicy.default)
            _head = []
            _subj = str(_msg.get("Subject", "") or "").strip()
            if _subj:
                _head.append("Subject: " + _subj)
            for _k in ("From", "To", "Cc", "Date"):
                _v = str(_msg.get(_k, "") or "").strip()
                if _v:
                    _head.append("%s: %s" % (_k, _v))
            _body, _atts = "", []
            if _msg.is_multipart():
                for _part in _msg.walk():
                    if (_part.get_content_disposition() or "") == "attachment":
                        _atts.append(str(_part.get_filename() or "unnamed"))
                        continue
                    if _part.get_content_type() == "text/plain" and not _body:
                        try:
                            _body = _part.get_content().strip()
                        except Exception:
                            pass
                if not _body:
                    for _part in _msg.walk():
                        if _part.get_content_type() == "text/html":
                            try:
                                _h = _part.get_content()
                            except Exception:
                                continue
                            try:
                                from bs4 import BeautifulSoup as _Soup2
                                _s2 = _Soup2(_h, "html.parser")
                                for _t2 in _s2(["script", "style"]):
                                    _t2.decompose()
                                _body = _s2.get_text(separator="\n").strip()
                            except Exception:
                                _body = re.sub(r"<[^>]+>", " ", _h).strip()
                            if _body:
                                break
            else:
                try:
                    _c = _msg.get_content()
                    _body = (_c if isinstance(_c, str) else "").strip()
                except Exception:
                    _body = ""
            _segs = []
            if _head:
                _segs.append("\n".join(_head))
            if _body:
                _segs.append(_body)
            if _atts:
                _segs.append("Attachments: " + "、".join(_atts))
            text = "\n\n".join(_segs)
        else:  # .pdf
            from pypdf import PdfReader
            reader = PdfReader(path)
            chunks = []
            for i, page in enumerate(reader.pages, 1):
                try:
                    t = page.extract_text() or ""
                except Exception:
                    t = ""
                if t.strip():
                    chunks.append("[Page %d]\n%s" % (i, t.strip()))
            text = "\n\n".join(chunks)
    except Exception as e:
        return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}
    text = (text or "").strip()
    if not text:
        return {"ok": False, "error": "抽不到文字（可能是掃描圖片型 PDF/ppt）"}
    truncated = len(text) > MAX_TEXT_CHARS
    return {"ok": True, "text": text[:MAX_TEXT_CHARS],
            "chars": len(text), "truncated": truncated}


def FnChatSummarize(text: str, source_names: List[str]) -> Dict[str, Any]:
    """One OpenAI-compatible chat call through the app-wide active LLM."""
    from .llm_provider import FnResolveActiveLLM
    tgt = FnResolveActiveLLM()
    if not tgt.get("ok"):
        return {"ok": False, "error": tgt.get("error") or "no LLM provider configured"}
    import json
    import urllib.request
    names = "、".join(source_names) or "文件"
    from .prompts import FnGetPrompt
    cfg = FnGetPrompt("ingest_text")
    template = cfg.get("template") or ""
    prompt = template.format(source_names=names, text=text)
    system_msg = cfg.get("system") or "你是繁體中文技術文件摘要助手，只輸出 Markdown。"
    payload = json.dumps({
        "model": tgt["model"],
        "messages": [
            {"role": "system", "content": system_msg},
            {"role": "user", "content": prompt},
        ],
        # NOTE: no temperature — some models (e.g. gpt-6-luna) only accept
        # the default value 1 and 400 on anything else.
    }).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if tgt.get("key"):
        headers["Authorization"] = "Bearer " + tgt["key"]
    try:
        req = urllib.request.Request(tgt["base"] + "/chat/completions",
                                     data=payload, headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=180) as res:
            data = json.loads(res.read().decode("utf-8"))
        msg = (((data.get("choices") or [{}])[0].get("message") or {}).get("content") or "").strip()
        if not msg:
            return {"ok": False, "error": "LLM 回傳為空"}
        return {"ok": True, "summary": msg, "model": tgt["model"], "provider": tgt["provider"]}
    except Exception as e:
        return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}


def _FnPilToDataUrl(img: Any) -> Optional[str]:
    """Normalize a PIL image (RGB, bounded, JPEG) -> data URL. None if trivial."""
    try:
        from PIL import Image
        import base64
        import io
    except Exception:
        return None
    try:
        if img.mode in ("RGBA", "LA", "P"):
            bg = Image.new("RGB", img.size, (255, 255, 255))
            try:
                bg.paste(img, mask=img.split()[-1] if img.mode in ("RGBA", "LA") else None)
            except Exception:
                bg.paste(img)
            img = bg
        elif img.mode != "RGB":
            img = img.convert("RGB")
        w, h = img.size
        if min(w, h) < VISION_MIN_EDGE:
            return None
        s = min(1.0, VISION_MAX_DIM / max(w, h))
        if s < 1.0:
            img = img.resize((max(1, int(w * s)), max(1, int(h * s))))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=VISION_JPEG_Q)
        return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")
    except Exception:
        return None


def FnExtractImages(path: str) -> Dict[str, Any]:
    """Render/extract page images for vision fallback.

    pdf  -> pymupdf page renders (covers scans AND odd encodings).
    pptx -> embedded slide pictures (no renderer available server-side).
    Returns {ok, images: [{label, data_url}], note}.
    """
    from PIL import Image
    import io as _io
    ext = os.path.splitext(path or "")[1].lower()
    images: List[Dict[str, str]] = []
    # Standalone photos/scans go straight to eyes (no page rendering).
    if ext in IMAGE_EXTS:
        try:
            img = Image.open(path)
            url = _FnPilToDataUrl(img)
        except Exception as e:
            return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}
        if not url:
            return {"ok": False, "error": "圖片太小或無法辨識"}
        return {"ok": True, "images": [{"label": "[Image]", "data_url": url}], "note": ""}
    try:
        if ext == ".pdf":
            import fitz
            doc = fitz.open(path)
            n = doc.page_count
            idxs = list(range(n)) if n <= VISION_MAX_PAGES else [
                int(i * n / VISION_MAX_PAGES) for i in range(VISION_MAX_PAGES)]
            for k, i in enumerate(idxs):
                if len(images) >= VISION_MAX_PAGES:
                    break
                page = doc[i]
                pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                url = _FnPilToDataUrl(img)
                if url:
                    images.append({"label": "[Page %d]" % (i + 1), "data_url": url})
            doc.close()
            if n > VISION_MAX_PAGES:
                return {"ok": True, "images": images,
                        "note": "%d 頁僅取 %d 頁辨識" % (n, len(images))}
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)
            for i, slide in enumerate(prs.slides, 1):
                if len(images) >= VISION_MAX_PAGES:
                    break
                for shape in slide.shapes:
                    try:
                        blob = shape.image.blob
                    except Exception:
                        continue
                    try:
                        img = Image.open(_io.BytesIO(blob))
                    except Exception:
                        continue
                    url = _FnPilToDataUrl(img)
                    if url:
                        images.append({"label": "[Slide %d]" % i, "data_url": url})
        else:
            return {"ok": False, "error": "不支援圖片辨識的格式: %s" % ext}
    except Exception as e:
        return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}
    if not images:
        return {"ok": False, "error": "找不到可辨識的頁面圖片（可能是加密或純向量檔）"}
    return {"ok": True, "images": images, "note": ""}


def FnVisionSummarize(images: List[Dict[str, str]],
                      source_names: List[str]) -> Dict[str, Any]:
    """Summarize page/slide images through a vision-capable chat model."""
    from .llm_provider import FnResolveActiveLLM
    tgt = FnResolveActiveLLM()
    if not tgt.get("ok"):
        return {"ok": False, "error": tgt.get("error") or "no LLM provider configured"}
    import json
    import urllib.request
    names = "、".join(source_names) or "文件"
    from .prompts import FnGetPrompt
    cfg = FnGetPrompt("ingest_vision")
    template = cfg.get("template") or ""
    prompt = template.format(source_names=names)
    system_msg = cfg.get("system") or "你是繁體中文技術文件摘要助手，只輸出 Markdown。"
    parts: List[Dict[str, Any]] = [{"type": "text", "text": prompt}]
    for im in images:
        parts.append({"type": "image_url",
                      "image_url": {"url": im["data_url"]}})
    payload = json.dumps({
        "model": tgt["model"],
        "messages": [
            {"role": "system", "content": system_msg},
            {"role": "user", "content": parts},
        ],
        # NOTE: no temperature (see FnChatSummarize) and no max_tokens —
        # quirky proxies 400 on non-standard params.
    }).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if tgt.get("key"):
        headers["Authorization"] = "Bearer " + tgt["key"]
    try:
        req = urllib.request.Request(tgt["base"] + "/chat/completions",
                                     data=payload, headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=180) as res:
            data = json.loads(res.read().decode("utf-8"))
        msg = (((data.get("choices") or [{}])[0].get("message") or {}).get("content") or "").strip()
        if not msg:
            return {"ok": False, "error": "LLM 回傳為空"}
        return {"ok": True, "summary": msg, "model": tgt["model"], "provider": tgt["provider"]}
    except Exception as e:
        return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}


def FnMultimodalSummarize(text: str, images: List[Dict[str, str]],
                          source_names: List[str],
                          pasted_image_rels: Optional[List[str]] = None) -> Dict[str, Any]:
    """Summarize text + images simultaneously (e.g. mixed clipboard text + screenshots)."""
    from .llm_provider import FnResolveActiveLLM
    tgt = FnResolveActiveLLM()
    if not tgt.get("ok"):
        return {"ok": False, "error": tgt.get("error") or "no LLM provider configured"}
    import json
    import urllib.request
    names = "、".join(source_names) or "剪貼簿混合素材"

    img_hints = []
    if pasted_image_rels:
        for idx, rel in enumerate(pasted_image_rels, 1):
            img_hints.append(f"- 截圖附件 {idx}：`![截圖說明]({rel})`")
    hints_text = "\n".join(img_hints)

    prompt = f"""請研讀以下剪貼簿提供的文字筆記，以及所附帶的截圖影像，進行嚴謹、結構化的技術與業務摘要。
若摘要中有說明到各張截圖的關鍵畫面、架構或介面，請在該段落適當處以 Markdown 語法引用該圖片（例如：`![畫面說明](檔案名稱.png)`）。

【可引用的截圖附件清單】：
{hints_text if hints_text else "（無額外圖片附件）"}

【筆記文字內容】：
===== 開始 =====
{text}
===== 結束 =====

請輸出完整、排版優良的繁體中文 Markdown 摘要。第一行請給出精準的 H1 主標題（格式：`# 主標題`）。"""

    system_msg = "你是繁體中文技術文件摘要助手，擅長解讀混合文字與截圖，只輸出 Markdown。"
    parts: List[Dict[str, Any]] = [{"type": "text", "text": prompt}]
    for im in images:
        parts.append({"type": "image_url",
                      "image_url": {"url": im["data_url"]}})

    payload = json.dumps({
        "model": tgt["model"],
        "messages": [
            {"role": "system", "content": system_msg},
            {"role": "user", "content": parts},
        ],
    }).encode("utf-8")
    headers = {"Content-Type": "application/json"}
    if tgt.get("key"):
        headers["Authorization"] = "Bearer " + tgt["key"]
    try:
        req = urllib.request.Request(tgt["base"] + "/chat/completions",
                                     data=payload, headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=180) as res:
            data = json.loads(res.read().decode("utf-8"))
        msg = (((data.get("choices") or [{}])[0].get("message") or {}).get("content") or "").strip()
        if not msg:
            return {"ok": False, "error": "LLM 回傳為空"}
        return {"ok": True, "summary": msg, "model": tgt["model"], "provider": tgt["provider"]}
    except Exception as e:
        return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}


def FnFetchUrlContent(url: str) -> Dict[str, Any]:
    """Fetch URL, extract page title and clean article body text."""
    import urllib.request
    import urllib.error
    import html
    clean_url = (url or "").strip()
    if not clean_url.startswith(("http://", "https://")):
        return {"ok": False, "error": "網址格式不正確，必須以 http:// 或 https:// 開頭"}

    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language": "zh-TW,zh;q=0.9,en-US;q=0.8,en;q=0.7",
    }
    try:
        req = urllib.request.Request(clean_url, headers=headers)
        with urllib.request.urlopen(req, timeout=30) as resp:
            content_type = resp.headers.get("Content-Type", "").lower()
            raw_bytes = resp.read()
            charset = "utf-8"
            if "charset=" in content_type:
                charset = content_type.split("charset=")[-1].split(";")[0].strip()
            try:
                raw_html = raw_bytes.decode(charset, errors="replace")
            except Exception:
                raw_html = raw_bytes.decode("utf-8", errors="replace")

        # Extract title
        title_m = re.search(r"<title[^>]*>(.*?)</title>", raw_html, re.IGNORECASE | re.DOTALL)
        raw_title = html.unescape(title_m.group(1)).strip() if title_m else ""
        raw_title = re.sub(r"[\r\n\t]+", " ", raw_title).strip()
        if not raw_title:
            raw_title = clean_url.split("/")[-1].split("?")[0] or "網頁文件"

        # Clean HTML scripts, styles, svg, noscript, header, footer, nav
        body_html = re.sub(r"<(script|style|svg|noscript|nav|header|footer)[^>]*>.*?</\1>", " ", raw_html, flags=re.IGNORECASE | re.DOTALL)
        # Convert paragraph/headings/breaks to newlines
        body_html = re.sub(r"</?(?:p|div|h\d|li|br|tr)[^>]*>", "\n", body_html, flags=re.IGNORECASE)
        # Strip remaining tags
        text = re.sub(r"<[^>]+>", " ", body_html)
        text = html.unescape(text)
        # Collapse whitespace
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.split("\n")]
        lines = [line for line in lines if line]
        clean_text = "\n\n".join(lines)

        if len(clean_text) < 30:
            return {"ok": False, "error": "抓取網頁成功，但未解析出足夠的內文"}

        truncated = len(clean_text) > MAX_TEXT_CHARS
        return {
            "ok": True,
            "url": clean_url,
            "title": raw_title,
            "text": clean_text[:MAX_TEXT_CHARS],
            "chars": len(clean_text),
            "truncated": truncated
        }
    except Exception as e:
        return {"ok": False, "error": f"網址抓取失敗: {e}"}


def FnLinkifyConcepts(text: str, official_map: Dict[str, str],
                     md_dir: str, tags_dir: str, *,
                     case_insensitive: bool = False,
                     skip_lines=None) -> str:
    """Wrap known concept mentions as physical md links to tag pages.

    official_map: {surface form -> canonical name} (official concepts only).
    Link target is the tag page path relative to the md's own directory, so
    the link works in any markdown renderer (Obsidian / VS Code preview).
    Skips fenced blocks, inline code, existing [...](...) links and
    legacy [[...]] brackets. Longest surface first.
    """
    if not text or not official_map:
        return text
    surfaces = sorted(
        (s for s in official_map.keys()
         if s and not any(c in s for c in "]()")),
        key=len, reverse=True)
    if not surfaces:
        return text
    skip_pat = re.compile(r"`[^`\n]+`|\[[^\]\n]*\]\([^)\n]*\)|\[\[[^\]\n]+\]\]")
    out_lines = []
    in_fence = False
    skip = skip_lines or set()
    for idx, line in enumerate(text.splitlines()):
        s = line.strip()
        if s.startswith("```") or s.startswith("~~~"):
            in_fence = not in_fence
            out_lines.append(line)
            continue
        if in_fence:
            out_lines.append(line)
            continue
        if idx in skip or s.startswith("#"):
            # Caller-excluded lines or Markdown headings: verbatim.
            out_lines.append(line)
            continue
        for surf in surfaces:
            canon = official_map[surf]
            tag_rel = os.path.relpath(
                os.path.join(tags_dir, canon + ".md"), md_dir).replace("\\", "/")
            spans = [(m.start(), m.end()) for m in skip_pat.finditer(line)]

            def _in_skip(a, b, _spans=spans):
                return any(a < e and b > st for st, e in _spans)

            if any(ord(c) > 127 for c in surf):
                pat = re.compile(re.escape(surf))
            else:
                # ASCII-word boundary ONLY: CJK neighbours are \w too and
                # must NOT block (e.g. 字CrawlNewA字 still linkifies).
                # Pure-ASCII behaviour is unchanged. case_insensitive is
                # opt-in (staged apply: evidence surface varies in case).
                flags = re.IGNORECASE if case_insensitive else 0
                pat = re.compile(r"(?<![A-Za-z0-9_])" + re.escape(surf) + r"(?![A-Za-z0-9_])",
                                 flags)
            res = []
            last = 0
            hit = False
            for m in pat.finditer(line):
                if _in_skip(m.start(), m.end()):
                    continue
                res.append(line[last:m.start()])
                res.append("[%s](%s)" % (m.group(0), tag_rel))
                last = m.end()
                hit = True
            if hit:
                line = "".join(res) + line[last:]
        out_lines.append(line)
    return "\n".join(out_lines)


def FnUniqueMdPath(target_dir: str, filename: str) -> str:
    """Collision-safe md path inside target_dir (adds -2, -3, ...)."""
    name = FnSanitizeFilename(filename)
    if not name.lower().endswith(".md"):
        name += ".md"
    cand = os.path.join(target_dir, name)
    if not os.path.exists(cand):
        return cand
    stem = name[:-3]
    i = 2
    while True:
        cand = os.path.join(target_dir, "%s-%d.md" % (stem, i))
        if not os.path.exists(cand):
            return cand
        i += 1


def FnBuildMd(title: str, source_rels: List[str], summary: str,
              model: str = "", source_urls: Optional[List[str]] = None) -> str:
    """Assemble the final md. Handles source files, URLs, and image attachments in the header."""
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    lines = ["# %s" % (title or "文件摘要"), ""]
    
    if source_urls:
        lines.append("> 來源網址：")
        for u in source_urls:
            lines.append("> - [%s](%s)" % (u, u))
        if source_rels:
            lines.append("> 附加檔案：")
            for r in source_rels:
                lines.append("> - `%s`" % r)
    else:
        lines.append("> 來源檔案（相對路徑）：")
        for r in source_rels:
            lines.append("> - `%s`" % r)
            
    lines.append(">")
    lines.append("> 匯入時間：%s%s" % (stamp, ("　生成模型：%s" % model) if model else ""))
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append((summary or "").strip())
    lines.append("")
    return "\n".join(lines)


def FnIngestCompoundTagging(summary_text: str, repo_path: str) -> Dict[str, Any]:
    """Pass-2 internal tagging pipeline (Compounding Ingest).

    Inspects the summary markdown against official concepts and candidate mining.
    Returns {ok, matched_concepts: [{name, evidence}], discovered_candidates: [{term, category, reason, aliases}]}.
    """
    from .llm_provider import FnResolveActiveLLM
    tgt = FnResolveActiveLLM()
    if not tgt.get("ok"):
        return {"ok": False, "error": tgt.get("error") or "no LLM provider configured"}

    import json
    import sqlite3
    import urllib.request
    from .db import get_db_path

    # 1. Fetch official concepts and blacklisted terms from db
    official_concepts = []
    blacklist_terms = []
    db_p = None
    try:
        db_p = get_db_path(repo_path)
    except Exception:
        db_p = None

    if db_p and os.path.isfile(db_p):
        try:
            con = sqlite3.connect(db_p)
            try:
                for r in con.execute("SELECT name, aliases, category FROM concepts WHERE status <> 'candidate'"):
                    c_name = r[0]
                    c_aliases = r[1] or ""
                    c_cat = r[2] or ""
                    alias_str = f" (別名: {c_aliases})" if c_aliases else ""
                    cat_str = f" [{c_cat}]" if c_cat else ""
                    official_concepts.append(f"- {c_name}{alias_str}{cat_str}")
                for r in con.execute("SELECT term FROM rejected_terms"):
                    if r[0]:
                        blacklist_terms.append(r[0])
            finally:
                con.close()
        except Exception as e:
            pass

    official_list_str = "\n".join(official_concepts) if official_concepts else "(尚無官方概念)"
    blacklist_str = "、".join(blacklist_terms) if blacklist_terms else "無"

    system_msg = (
        "你是知識庫拓撲架構師。你的唯一天職是依據技術摘要，比對現有官方概念並發掘新領域術語。"
        "嚴格只輸出單一合法的 JSON 物件，絕不輸出任何 Markdown 代碼塊或額外解釋。"
    )

    user_prompt = f"""【現有官方概念清單】：
{official_list_str}

【黑名單排除詞彙（絕對不可再提）】：
{blacklist_str}

【輸入技術文件摘要內容】：
===== 摘要開始 =====
{summary_text[:8000]}
===== 摘要結束 =====

請仔細研讀上述摘要，精準輸出以下 JSON 格式：
{{
  "matched_concepts": [
    {{"name": "吻合的官方概念名", "evidence": "摘要中提及此概念的佐證段落（15字內）"}}
  ],
  "discovered_candidates": [
    {{"term": "高價值新領域專有名詞（設備、代號、工藝，嚴禁通用詞）", "category": "建議分類", "reason": "收錄原因（10字內）", "aliases": "可能的別名或縮寫"}}
  ]
}}
無吻合或無新詞時請給空陣列 []。只輸出純 JSON。"""

    payload = json.dumps({
        "model": tgt["model"],
        "messages": [
            {"role": "system", "content": system_msg},
            {"role": "user", "content": user_prompt},
        ],
    }).encode("utf-8")

    headers = {"Content-Type": "application/json"}
    if tgt.get("key"):
        headers["Authorization"] = "Bearer " + tgt["key"]

    try:
        req = urllib.request.Request(tgt["base"] + "/chat/completions",
                                     data=payload, headers=headers, method="POST")
        with urllib.request.urlopen(req, timeout=90) as res:
            raw_data = json.loads(res.read().decode("utf-8"))
        msg = (((raw_data.get("choices") or [{}])[0].get("message") or {}).get("content") or "").strip()
        if not msg:
            return {"ok": False, "error": "LLM 打標回傳為空"}

        # Strip possible markdown code fences
        cleaned = re.sub(r"^```(?:json)?\s*", "", msg, flags=re.MULTILINE)
        cleaned = re.sub(r"\s*```$", "", cleaned, flags=re.MULTILINE).strip()
        parsed = json.loads(cleaned)

        matched = parsed.get("matched_concepts") or []
        candidates = parsed.get("discovered_candidates") or []

        return {
            "ok": True,
            "matched_concepts": matched,
            "discovered_candidates": candidates,
            "model": tgt["model"]
        }
    except Exception as e:
        return {"ok": False, "error": "%s: %s" % (type(e).__name__, e)}